from http.server import BaseHTTPRequestHandler
from io import BytesIO
from datetime import datetime, date
import cgi
import json
import re

from openpyxl import load_workbook

SCHOOLS = {"RDXB": "Repton DXB", "RAB": "RAB", "FRY": "Fry", "ROSE": "Rose"}


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def as_number(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 20) + 1):
        values = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "DATE" in values and any(v == "TOTAL" for v in values):
            return r
    raise ValueError(f"Could not find the sales header row in sheet '{ws.title}'.")


def parse_source(file_bytes, filename):
    wb = load_workbook(BytesIO(file_bytes), data_only=False, read_only=True)
    ws = wb.active
    header_row = find_header_row(ws)
    headers = {norm(ws.cell(header_row, c).value): c for c in range(1, ws.max_column + 1)}

    date_col = headers.get("DATE")
    total_col = headers.get("TOTAL")
    exchange_col = next((c for h, c in headers.items() if "EXCHANGE" in h), None)
    campus_col = headers.get("CAMPUS")
    if not date_col or not total_col or not exchange_col:
        raise ValueError(f"{filename}: Date, TOTAL or Exchange column was not found.")

    # TOTAL is a formula in Zona reports. Sum the numeric transaction columns that feed it.
    transaction_cols = list(range(4, total_col))
    rows = []
    campuses = set()

    fallback_campus = next((s for s in SCHOOLS if s in norm(filename)), None)
    if not fallback_campus and norm(ws.title) in SCHOOLS:
        fallback_campus = norm(ws.title)

    for r in range(header_row + 1, ws.max_row + 1):
        d = as_date(ws.cell(r, date_col).value)
        if not d:
            continue
        campus = norm(ws.cell(r, campus_col).value) if campus_col else fallback_campus
        campus = campus if campus in SCHOOLS else fallback_campus
        if campus not in SCHOOLS:
            continue
        campuses.add(campus)
        total = sum(as_number(ws.cell(r, c).value) for c in transaction_cols)
        exchange = as_number(ws.cell(r, exchange_col).value)
        rows.append((d, campus, total, exchange))

    if not rows:
        raise ValueError(f"{filename}: no usable dated sales rows were found.")
    return rows, campuses


def choose_rows(parsed):
    combined = [(name, rows, campuses) for name, rows, campuses in parsed if set(SCHOOLS).issubset(campuses)]
    if combined:
        # Prefer the largest complete combined report and never mix it with individual reports.
        return max(combined, key=lambda item: len(item[1]))[1]

    by_school = {}
    for name, rows, campuses in parsed:
        if len(campuses) == 1:
            school = next(iter(campuses))
            by_school[school] = rows
    missing = [s for s in SCHOOLS if s not in by_school]
    if missing:
        raise ValueError("Missing source report(s): " + ", ".join(missing) + ". Upload a complete combined report or all four school reports.")
    return [row for school in SCHOOLS for row in by_school[school]]


def find_target_columns(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = {norm(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1)}
        school_cols = {}
        for code, label in SCHOOLS.items():
            aliases = {norm(label), code}
            school_cols[code] = next((c for h, c in vals.items() if h in aliases), None)
        exchange = next((c for h, c in vals.items() if h == "EXCHANGE"), None)
        date_cols = [c for h, c in vals.items() if h == "DATE"]
        if exchange and all(school_cols.values()) and date_cols:
            return r, school_cols, exchange, max(date_cols)
    raise ValueError(f"Could not identify target columns in '{ws.title}'.")


def update_target(target_bytes, rows):
    wb = load_workbook(BytesIO(target_bytes), data_only=False)
    sales = {}
    exchanges = {}
    for d, campus, total, exchange in rows:
        sales[(d, campus)] = sales.get((d, campus), 0) + total
        exchanges[d] = exchanges.get(d, 0) + exchange

    dates = sorted({d for d, _, _, _ in rows})
    for d in dates:
        sheet_name = d.strftime("%b-%y")
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Target workbook has no sheet named '{sheet_name}' for {d:%d %b %Y}.")
        ws = wb[sheet_name]
        header_row, school_cols, exchange_col, date_col = find_target_columns(ws)
        target_row = None
        for r in range(header_row + 1, ws.max_row + 1):
            if as_date(ws.cell(r, date_col).value) == d:
                target_row = r
                break
        if not target_row:
            raise ValueError(f"Could not find {d:%d %b %Y} in {sheet_name}.")

        for campus, col in school_cols.items():
            ws.cell(target_row, col).value = sales.get((d, campus), 0)
        ws.cell(target_row, exchange_col).value = exchanges.get(d, 0)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={
                "REQUEST_METHOD": "POST",
                "CONTENT_TYPE": self.headers.get("Content-Type"),
            })
            target = form["target"] if "target" in form else None
            source_field = form["sources"] if "sources" in form else None
            if target is None or source_field is None:
                raise ValueError("Upload the target tracker and supplier source report(s).")
            sources = source_field if isinstance(source_field, list) else [source_field]
            parsed = []
            for item in sources:
                data = item.file.read()
                rows, campuses = parse_source(data, item.filename or "source.xlsx")
                parsed.append((item.filename or "source.xlsx", rows, campuses))
            selected_rows = choose_rows(parsed)
            result = update_target(target.file.read(), selected_rows)

            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", 'attachment; filename="Uniform Daily Sales Tracker - Updated.xlsx"')
            self.send_header("Content-Length", str(len(result)))
            self.end_headers()
            self.wfile.write(result)
        except Exception as exc:
            body = json.dumps({"detail": str(exc)}).encode("utf-8")
            self.send_response(400)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
