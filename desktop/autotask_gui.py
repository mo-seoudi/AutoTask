import os
import re
import threading
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

SCHOOLS = {"RDXB": "Repton DXB", "RAB": "RAB", "FRY": "Fry", "ROSE": "Rose"}


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def as_number(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def as_date(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    if isinstance(value, str):
        for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try: return datetime.strptime(value.strip(), fmt).date()
            except ValueError: pass
    return None


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 20) + 1):
        values = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "DATE" in values and "TOTAL" in values: return r
    raise ValueError(f"Could not find the sales header row in sheet '{ws.title}'.")


def parse_source(path):
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb.active
    header_row = find_header_row(ws)
    headers = {norm(ws.cell(header_row, c).value): c for c in range(1, ws.max_column + 1)}
    date_col, total_col = headers.get("DATE"), headers.get("TOTAL")
    exchange_col = next((c for h, c in headers.items() if "EXCHANGE" in h), None)
    campus_col = headers.get("CAMPUS")
    if not date_col or not total_col or not exchange_col:
        raise ValueError(f"{Path(path).name}: Date, TOTAL or Exchange column was not found.")
    transaction_cols = list(range(4, total_col))
    rows, campuses = [], set()
    fallback = next((s for s in SCHOOLS if s in norm(Path(path).name)), None)
    if not fallback and norm(ws.title) in SCHOOLS: fallback = norm(ws.title)
    for r in range(header_row + 1, ws.max_row + 1):
        d = as_date(ws.cell(r, date_col).value)
        if not d: continue
        campus = norm(ws.cell(r, campus_col).value) if campus_col else fallback
        campus = campus if campus in SCHOOLS else fallback
        if campus not in SCHOOLS: continue
        campuses.add(campus)
        total = sum(as_number(ws.cell(r, c).value) for c in transaction_cols)
        exchange = as_number(ws.cell(r, exchange_col).value)
        rows.append((d, campus, total, exchange))
    if not rows: raise ValueError(f"{Path(path).name}: no usable dated sales rows were found.")
    return rows, campuses


def choose_rows(parsed):
    combined = [(p, rows, campuses) for p, rows, campuses in parsed if set(SCHOOLS).issubset(campuses)]
    if combined: return max(combined, key=lambda item: len(item[1]))[1]
    by_school = {}
    for _, rows, campuses in parsed:
        if len(campuses) == 1: by_school[next(iter(campuses))] = rows
    missing = [s for s in SCHOOLS if s not in by_school]
    if missing: raise ValueError("Missing source report(s): " + ", ".join(missing))
    return [row for school in SCHOOLS for row in by_school[school]]


def find_target_columns(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = [(norm(ws.cell(r, c).value), c) for c in range(1, ws.max_column + 1)]
        school_cols = {code: next((c for h, c in vals if h in {norm(label), code}), None) for code, label in SCHOOLS.items()}
        exchange = next((c for h, c in vals if h == "EXCHANGE"), None)
        date_cols = [c for h, c in vals if h == "DATE"]
        if exchange and all(school_cols.values()) and date_cols:
            def score(col): return sum(as_date(ws.cell(rr, col).value) is not None for rr in range(r + 1, ws.max_row + 1))
            return r, school_cols, exchange, max(date_cols, key=score)
    raise ValueError(f"Could not identify target columns in '{ws.title}'.")


def update_target(target_path, rows, output_path):
    wb = load_workbook(target_path, data_only=False)
    sales, exchanges = {}, {}
    for d, campus, total, exchange in rows:
        sales[(d, campus)] = sales.get((d, campus), 0) + total
        exchanges[d] = exchanges.get(d, 0) + exchange
    for d in sorted({x[0] for x in rows}):
        sheet = d.strftime("%b-%y")
        if sheet not in wb.sheetnames: raise ValueError(f"Target workbook has no sheet named '{sheet}'.")
        ws = wb[sheet]
        header, school_cols, exchange_col, date_col = find_target_columns(ws)
        target_row = next((r for r in range(header + 1, ws.max_row + 1) if as_date(ws.cell(r, date_col).value) == d), None)
        if not target_row: raise ValueError(f"Could not find {d:%d %b %Y} in {sheet}.")
        for campus, col in school_cols.items(): ws.cell(target_row, col).value = sales.get((d, campus), 0)
        ws.cell(target_row, exchange_col).value = exchanges.get(d, 0)
    wb.save(output_path)


class AutoTaskGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("AutoTask — Zona Daily Sales")
        self.geometry("760x520")
        self.minsize(700, 480)
        self.configure(bg="#f5f7fa")
        self.target = tk.StringVar()
        self.sources = []
        self.source_text = tk.StringVar(value="No files selected")
        self.status = tk.StringVar(value="Ready")
        self._build()

    def _build(self):
        style = ttk.Style(self)
        style.theme_use("vista" if "vista" in style.theme_names() else "clam")
        style.configure("Title.TLabel", font=("Segoe UI", 19, "bold"), background="#f5f7fa", foreground="#172033")
        style.configure("Sub.TLabel", font=("Segoe UI", 9), background="#f5f7fa", foreground="#6b7280")
        style.configure("Card.TFrame", background="white")
        style.configure("Card.TLabel", background="white", foreground="#172033", font=("Segoe UI", 10, "bold"))
        style.configure("Hint.TLabel", background="white", foreground="#7b8493", font=("Segoe UI", 9))

        outer = ttk.Frame(self, padding=(42, 30, 42, 30)); outer.pack(fill="both", expand=True)
        ttk.Label(outer, text="AutoTask", style="Sub.TLabel").pack(anchor="w")
        ttk.Label(outer, text="Zona Daily Sales", style="Title.TLabel").pack(anchor="w", pady=(3, 2))
        ttk.Label(outer, text="Update the Uniform Daily Sales Tracker from Zona supplier reports.", style="Sub.TLabel").pack(anchor="w", pady=(0, 20))

        card = ttk.Frame(outer, style="Card.TFrame", padding=24); card.pack(fill="both", expand=True)
        ttk.Label(card, text="1   Daily Sales Tracker", style="Card.TLabel").pack(anchor="w")
        ttk.Label(card, text="Select the current master workbook.", style="Hint.TLabel").pack(anchor="w", pady=(3, 8))
        row1 = ttk.Frame(card, style="Card.TFrame"); row1.pack(fill="x")
        ttk.Entry(row1, textvariable=self.target, state="readonly").pack(side="left", fill="x", expand=True)
        ttk.Button(row1, text="Browse…", command=self.pick_target).pack(side="left", padx=(8, 0))

        ttk.Separator(card).pack(fill="x", pady=22)
        ttk.Label(card, text="2   Zona report(s)", style="Card.TLabel").pack(anchor="w")
        ttk.Label(card, text="Choose the combined report or all four school reports.", style="Hint.TLabel").pack(anchor="w", pady=(3, 8))
        row2 = ttk.Frame(card, style="Card.TFrame"); row2.pack(fill="x")
        ttk.Entry(row2, textvariable=self.source_text, state="readonly").pack(side="left", fill="x", expand=True)
        ttk.Button(row2, text="Browse…", command=self.pick_sources).pack(side="left", padx=(8, 0))

        ttk.Separator(card).pack(fill="x", pady=22)
        actions = ttk.Frame(card, style="Card.TFrame"); actions.pack(fill="x")
        self.run_btn = ttk.Button(actions, text="Process & Save", command=self.start_processing); self.run_btn.pack(side="left")
        self.progress = ttk.Progressbar(actions, mode="indeterminate", length=120); self.progress.pack(side="left", padx=12)
        ttk.Label(actions, textvariable=self.status, style="Hint.TLabel").pack(side="left")

    def pick_target(self):
        p = filedialog.askopenfilename(title="Select Uniform Daily Sales Tracker", filetypes=[("Excel workbook", "*.xlsx")])
        if p: self.target.set(p)

    def pick_sources(self):
        paths = filedialog.askopenfilenames(title="Select Zona report(s)", filetypes=[("Excel workbook", "*.xlsx")])
        if paths:
            self.sources = list(paths)
            self.source_text.set(f"{len(paths)} file(s): " + ", ".join(Path(p).name for p in paths))

    def start_processing(self):
        if not self.target.get() or not self.sources:
            messagebox.showwarning("Missing files", "Select the tracker and Zona report(s) first.")
            return
        default = Path(self.target.get()).with_name(Path(self.target.get()).stem + " - Updated.xlsx")
        output = filedialog.asksaveasfilename(title="Save updated tracker", initialdir=default.parent, initialfile=default.name, defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not output: return
        self.run_btn.state(["disabled"]); self.progress.start(10); self.status.set("Processing…")
        threading.Thread(target=self._process, args=(output,), daemon=True).start()

    def _process(self, output):
        try:
            parsed = [(p, *parse_source(p)) for p in self.sources]
            rows = choose_rows(parsed)
            update_target(self.target.get(), rows, output)
            self.after(0, lambda: self._finish(output, None))
        except Exception as exc:
            self.after(0, lambda e=str(exc): self._finish(None, e))

    def _finish(self, output, error):
        self.progress.stop(); self.run_btn.state(["!disabled"])
        if error:
            self.status.set("Failed")
            messagebox.showerror("Processing failed", error)
        else:
            self.status.set("Completed")
            if messagebox.askyesno("Completed", f"Updated tracker saved successfully.\n\n{output}\n\nOpen the file now?"):
                os.startfile(output)


if __name__ == "__main__":
    AutoTaskGUI().mainloop()
